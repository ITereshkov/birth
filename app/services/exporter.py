from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook

from app.models.entities import TxType
from app.repositories.user_db import UserRepository
from app.services.reports import ReportService


class ExportService:
    def __init__(self, report_service: ReportService, template_path: Path | None = None) -> None:
        self.report_service = report_service
        self.template_path = template_path

    def export_xlsx(self, repo: UserRepository, start: date, end: date, out_path: Path) -> Path:
        summary = self.report_service.summarize(repo, start, end)
        if self.template_path and self.template_path.exists():
            wb = load_workbook(self.template_path)
            self._fill_template(wb, summary)
            out_path.parent.mkdir(parents=True, exist_ok=True)
            wb.save(out_path)
            return out_path

        wb = Workbook()
        ws_income = wb.active
        ws_income.title = "Доходы"
        ws_expense = wb.create_sheet("Расходы")
        ws_all = wb.create_sheet("Транзакции")
        ws_summary = wb.create_sheet("Сводка")

        headers = ["Дата", "Сумма", "Категория", "Тип", "Комментарий"]
        for ws in (ws_income, ws_expense, ws_all):
            ws.append(headers)

        for row in summary["transactions"]:
            values = [
                row["occurred_date"],
                row["amount"],
                row["category"],
                row["type"],
                row["comment"] or "",
            ]
            ws_all.append(values)
            if row["type"] == TxType.INCOME.value:
                ws_income.append(values)
            else:
                ws_expense.append(values)

        ws_summary.append(["Показатель", "Значение"])
        ws_summary.append(["Доходы", summary["income"]])
        ws_summary.append(["Расходы", summary["expense"]])
        ws_summary.append(["Разница", summary["difference"]])
        ws_summary.append([])
        ws_summary.append(["Топ-расходы", "Сумма"])
        for category, amount in summary["top_expenses"]:
            ws_summary.append([category, amount])

        out_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(out_path)
        return out_path

    def _fill_template(self, wb, summary: dict) -> None:
        txs = summary["transactions"]
        income_map: dict[str, float] = {}
        expense_map: dict[str, float] = {}
        for row in txs:
            target = income_map if row["type"] == TxType.INCOME.value else expense_map
            key = str(row["category"]).strip().lower()
            target[key] = target.get(key, 0.0) + float(row["amount"])

        if "Доходы" in wb.sheetnames:
            self._fill_category_amount_sheet(wb["Доходы"], income_map)
        if "Расходы" in wb.sheetnames:
            self._fill_category_amount_sheet(wb["Расходы"], expense_map)
        if "Баланс" in wb.sheetnames:
            self._fill_balance_sheet(wb["Баланс"], income_map, expense_map)

    def _fill_category_amount_sheet(self, ws, category_amount: dict[str, float]) -> None:
        for row in range(1, ws.max_row + 1):
            cat_cell = ws.cell(row=row, column=2).value
            if not cat_cell:
                continue
            cat_name = str(cat_cell).strip().lower()
            if cat_name.startswith("итого") or cat_name.startswith("всего"):
                continue
            amount = category_amount.get(cat_name)
            if amount is not None:
                ws.cell(row=row, column=3).value = round(amount, 2)

    def _fill_balance_sheet(self, ws, income_map: dict[str, float], expense_map: dict[str, float]) -> None:
        # Левая таблица доходов (категория в col B, сумма в col C)
        # Правая таблица расходов (категория в col G, сумма в col I)
        for row in range(1, ws.max_row + 1):
            lcat = ws.cell(row=row, column=2).value
            if lcat:
                name = str(lcat).strip().lower()
                if not name.startswith("итого"):
                    val = income_map.get(name)
                    if val is not None:
                        ws.cell(row=row, column=3).value = round(val, 2)

            rcat = ws.cell(row=row, column=7).value
            if rcat:
                name = str(rcat).strip().lower()
                if not name.startswith("итого"):
                    val = expense_map.get(name)
                    if val is not None:
                        ws.cell(row=row, column=9).value = round(val, 2)
